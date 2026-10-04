#!/usr/bin/env python3
"""
Convert E-Invoice Markdown documentation into beautifully styled Word (.docx) files.
Specifically formatted for submission to the Taiwan Ministry of Finance (MOF)
Fiscal Information Agency and National Taxation Bureau of Taipei (Zhongzheng Branch).
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = "/invoice/EINVTurnkey/docs"

# Typography & Palette
FONT_CHINESE = "Microsoft JhengHei"
FONT_ENGLISH = "Calibri"
COLOR_PRIMARY = RGBColor(0x1F, 0x29, 0x37)   # Charcoal text
COLOR_TITLE = RGBColor(0x11, 0x18, 0x27)     # Deep dark charcoal
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # Official Navy
COLOR_HEADER_BG = "F3F4F6"                   # Very light gray for table headers
COLOR_ALT_BG = "FAFAFA"                      # Alternate row
BORDER_COLOR = "9CA3AF"                      # Table border gray
BORDER_LIGHT = "D1D5DB"                      # Table inner border light gray

def set_run_font(run, font_name=FONT_CHINESE, size_pt=10.5, bold=False, italic=False, color=None):
    run.font.name = font_name
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set inner cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'''
        <w:tcMar %s>
            <w:top w:w="%d" w:type="dxa"/>
            <w:bottom w:w="%d" w:type="dxa"/>
            <w:left w:w="%d" w:type="dxa"/>
            <w:right w:w="%d" w:type="dxa"/>
        </w:tcMar>
    ''' % (nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill_hex))
    cell._tc.get_or_add_tcPr().append(shd)

def apply_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="%s"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="%s"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="%s"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="%s"/>
        </w:tblBorders>
    ''' % (nsdecls('w'), BORDER_COLOR, BORDER_COLOR, BORDER_COLOR, BORDER_COLOR, BORDER_LIGHT, BORDER_LIGHT))
    tblPr.append(borders)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))

def parse_formatted_text(p, text, base_size=10.5, default_bold=False, default_color=None):
    """Parses **bold**, *italic*, `code`, [x], [ ] into docx runs."""
    text = text.replace("- [x]", "☑").replace("- [ ]", "☐")
    text = text.replace("[x]", "☑").replace("[ ]", "☐")
    
    token_pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = token_pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            set_run_font(run, size_pt=base_size, bold=True, color=default_color or COLOR_PRIMARY)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            set_run_font(run, font_name="Consolas", size_pt=base_size * 0.95, bold=True, color=RGBColor(0x1F, 0x29, 0x37))
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            set_run_font(run, size_pt=base_size, italic=True, color=default_color or COLOR_PRIMARY)
        else:
            run = p.add_run(token)
            set_run_font(run, size_pt=base_size, bold=default_bold, color=default_color or COLOR_PRIMARY)

def add_seal_box(doc, has_rep_id=False):
    """Creates a dedicated official Taiwanese company & representative stamping table."""
    p_lead = doc.add_paragraph()
    p_lead.paragraph_format.space_before = Pt(14)
    p_lead.paragraph_format.space_after = Pt(4)
    run_lead = p_lead.add_run("【營業人及代表人簽章用印欄】")
    set_run_font(run_lead, size_pt=11, bold=True, color=COLOR_NAVY)
    
    seal_table = doc.add_table(rows=2, cols=2)
    seal_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(seal_table)
    
    for row in seal_table.rows:
        make_row_cant_split(row)
        row.cells[0].width = Inches(3.33)
        row.cells[1].width = Inches(3.33)
    
    # Row 0: Stamping Areas
    # Cell 0: Company Seal
    c0 = seal_table.rows[0].cells[0]
    set_cell_margins(c0, top=140, bottom=140, left=180, right=180)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run("營業人印鑑章（公司大章）\n")
    set_run_font(r0, size_pt=10.5, bold=True)
    
    p0_box = c0.add_paragraph()
    p0_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0_box.paragraph_format.space_before = Pt(24)
    p0_box.paragraph_format.space_after = Pt(24)
    r0_box = p0_box.add_run("（ 請 蓋 公 司 印 鑑 章 ）")
    set_run_font(r0_box, size_pt=11, italic=True, color=RGBColor(0x9C, 0xA3, 0xAF))
    
    # Cell 1: Representative Seal
    c1 = seal_table.rows[0].cells[1]
    set_cell_margins(c1, top=140, bottom=140, left=180, right=180)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("負責人印鑑章（代表人小章）\n")
    set_run_font(r1, size_pt=10.5, bold=True)
    
    p1_box = c1.add_paragraph()
    p1_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1_box.paragraph_format.space_before = Pt(24)
    p1_box.paragraph_format.space_after = Pt(24)
    r1_box = p1_box.add_run("（ 請 蓋 負 責 人 印 章 ）")
    set_run_font(r1_box, size_pt=11, italic=True, color=RGBColor(0x9C, 0xA3, 0xAF))
    
    # Row 1: Company & Representative Metadata details
    c_info0 = seal_table.rows[1].cells[0]
    set_cell_margins(c_info0, top=100, bottom=100, left=150, right=150)
    set_cell_background(c_info0, "FAFAFA")
    p_info0 = c_info0.paragraphs[0]
    p_info0.paragraph_format.line_spacing = 1.3
    p_info0.paragraph_format.space_after = Pt(2)
    lines0 = [
        "營業人名稱：奧銳有限公司",
        "統一編號：00015555",
        "稅籍編號：100213880",
        "營業登記地址：臺北市中正區漢口街1段45號10樓"
    ]
    for idx, l in enumerate(lines0):
        if idx > 0:
            p_info0 = c_info0.add_paragraph()
            p_info0.paragraph_format.line_spacing = 1.3
            p_info0.paragraph_format.space_after = Pt(2)
        parse_formatted_text(p_info0, l, base_size=9.5)
        
    c_info1 = seal_table.rows[1].cells[1]
    set_cell_margins(c_info1, top=100, bottom=100, left=150, right=150)
    set_cell_background(c_info1, "FAFAFA")
    p_info1 = c_info1.paragraphs[0]
    p_info1.paragraph_format.line_spacing = 1.3
    p_info1.paragraph_format.space_after = Pt(2)
    lines1 = [
        "負責人姓名：王世全",
        "身分證字號：L122116338",
        "聯絡電話：0903888022",
        "電子信箱：paul@wang.net"
    ]
    for idx, l in enumerate(lines1):
        if idx > 0:
            p_info1 = c_info1.add_paragraph()
            p_info1.paragraph_format.line_spacing = 1.3
            p_info1.paragraph_format.space_after = Pt(2)
        parse_formatted_text(p_info1, l, base_size=9.5)

    # Date line below table
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(14)
    p_date.paragraph_format.space_after = Pt(8)
    r_date = p_date.add_run("中華民國 115 年 09 月 24 日")
    set_run_font(r_date, size_pt=11, bold=True)

def get_col_widths(col_count):
    """Calculates column widths to fit 6.67 inches total width."""
    if col_count == 3:
        # Col 0: 1.8 in, Col 1: 3.2 in, Col 2: 1.67 in
        return [Inches(1.8), Inches(3.2), Inches(1.67)]
    elif col_count == 5:
        # Col 0: 0.5 in, Col 1: 1.2 in, Col 2: 2.2 in, Col 3: 0.8 in, Col 4: 1.97 in
        return [Inches(0.5), Inches(1.2), Inches(2.2), Inches(0.8), Inches(1.97)]
    elif col_count == 2:
        return [Inches(2.5), Inches(4.17)]
    elif col_count == 4:
        return [Inches(1.2), Inches(2.3), Inches(1.2), Inches(1.97)]
    else:
        # Equal division
        w = Inches(6.67 / col_count)
        return [w] * col_count

def convert_md_file(md_path, docx_path):
    print(f"Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = Document()
    
    # Configure A4 page margins (top/bottom: 0.8 in, left/right: 0.8 in)
    section = doc.sections[0]
    section.page_width = Inches(8.27)    # 210mm
    section.page_height = Inches(11.69)  # 297mm
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    # Base Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = FONT_CHINESE
    style_normal._element.rPr.get_or_add_rFonts().set(qn('w:eastAsia'), FONT_CHINESE)
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_PRIMARY
    
    i = 0
    in_code_block = False
    code_lines = []
    seal_box_added = False
    
    # Check if this document requires a seal box
    is_statutory_form = any(k in os.path.basename(md_path) for k in ["01_", "02_", "03_", "04_", "06_"])
    has_rep_id = ("02_" in os.path.basename(md_path))
    
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        
        # Check for code blocks ``` ... ```
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.2)
                p_text = "\n".join(code_lines)
                run = p.add_run(p_text)
                set_run_font(run, font_name="Consolas", size_pt=9.0, color=RGBColor(0x37, 0x41, 0x51))
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Skip empty lines
        if not stripped:
            i += 1
            continue
            
        # Skip horizontal rules
        if stripped in ["---", "***", "___"]:
            i += 1
            continue
            
        # Title (# )
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title_text)
            set_run_font(run, size_pt=18, bold=True, color=COLOR_TITLE)
            i += 1
            continue
            
        # Heading 2 (## )
        if stripped.startswith("## "):
            h2_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h2_text)
            set_run_font(run, size_pt=13, bold=True, color=COLOR_NAVY)
            i += 1
            continue
            
        # Heading 3 (### )
        if stripped.startswith("### "):
            h3_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h3_text)
            set_run_font(run, size_pt=11.5, bold=True, color=COLOR_TITLE)
            i += 1
            continue
            
        # Metadata / Blockquote (> )
        if stripped.startswith(">"):
            bq_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                raw_bq = lines[i].strip()[1:].strip()
                if raw_bq.startswith("[!"):
                    i += 1
                    continue
                bq_lines.append(raw_bq)
                i += 1
            
            p_meta = doc.add_paragraph()
            p_meta.paragraph_format.space_before = Pt(4)
            p_meta.paragraph_format.space_after = Pt(8)
            p_meta.paragraph_format.left_indent = Inches(0.2)
            p_meta.paragraph_format.line_spacing = 1.25
            
            for b_idx, b_line in enumerate(bq_lines):
                if b_idx > 0:
                    run_sep = p_meta.add_run("\n")
                    set_run_font(run_sep, size_pt=10)
                parse_formatted_text(p_meta, b_line, base_size=10, default_color=RGBColor(0x37, 0x41, 0x51))
            continue
            
        # Table detection (| ... |)
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                header_cells = [c.strip() for c in table_lines[0].strip('|').split('|')]
                col_count = len(header_cells)
                col_widths = get_col_widths(col_count)
                
                data_rows = []
                for row_line in table_lines[1:]:
                    cells = [c.strip() for c in row_line.strip('|').split('|')]
                    if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                        continue
                    if len(cells) == col_count:
                        data_rows.append(cells)
                    elif len(cells) < col_count:
                        cells.extend([""] * (col_count - len(cells)))
                        data_rows.append(cells)
                    else:
                        data_rows.append(cells[:col_count])
                        
                table = doc.add_table(rows=1 + len(data_rows), cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                apply_table_borders(table)
                
                # Header row
                hdr_row = table.rows[0]
                make_row_header(hdr_row)
                make_row_cant_split(hdr_row)
                for c_idx, h_text in enumerate(header_cells):
                    cell = hdr_row.cells[c_idx]
                    cell.width = col_widths[c_idx]
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                    set_cell_background(cell, COLOR_HEADER_BG)
                    p_hdr = cell.paragraphs[0]
                    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_hdr.paragraph_format.space_after = Pt(0)
                    parse_formatted_text(p_hdr, h_text, base_size=10, default_bold=True, default_color=COLOR_TITLE)
                    
                # Data rows
                for r_idx, row_data in enumerate(data_rows):
                    row = table.rows[1 + r_idx]
                    make_row_cant_split(row)
                    for c_idx, val in enumerate(row_data):
                        cell = row.cells[c_idx]
                        cell.width = col_widths[c_idx]
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        if r_idx % 2 == 1:
                            set_cell_background(cell, COLOR_ALT_BG)
                            
                        sub_lines = re.split(r'<br\s*/?>', val, flags=re.IGNORECASE)
                        cell_p = cell.paragraphs[0]
                        cell_p.paragraph_format.line_spacing = 1.2
                        cell_p.paragraph_format.space_after = Pt(1)
                        
                        clean_first = sub_lines[0].strip()
                        if clean_first.isdigit() or clean_first in ["符合", "☑ 符合", "年配", "期配", "07", "08"]:
                            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif c_idx == 0 and col_count > 3:
                            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            
                        for s_idx, s_line in enumerate(sub_lines):
                            s_line = s_line.strip()
                            if not s_line:
                                continue
                            if s_idx > 0:
                                cell_p = cell.add_paragraph()
                                cell_p.paragraph_format.line_spacing = 1.2
                                cell_p.paragraph_format.space_after = Pt(1)
                            parse_formatted_text(cell_p, s_line, base_size=9.5)
                            
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(4)
                p_spacer.paragraph_format.space_after = Pt(4)
            continue
            
        # Lists / Checkbox items
        if stripped.startswith("- ") or stripped.startswith("* ") or re.match(r'^\d+\.\s', stripped):
            p_list = doc.add_paragraph()
            p_list.paragraph_format.space_before = Pt(2)
            p_list.paragraph_format.space_after = Pt(2)
            p_list.paragraph_format.left_indent = Inches(0.25)
            p_list.paragraph_format.line_spacing = 1.25
            
            content_text = stripped
            if stripped.startswith("- ") or stripped.startswith("* "):
                content_text = stripped[2:].strip()
            elif re.match(r'^\d+\.\s', stripped):
                content_text = stripped
                
            parse_formatted_text(p_list, content_text, base_size=10)
            i += 1
            continue
            
        # Closing & Seal section
        if is_statutory_form and not seal_box_added and ("此致" in stripped or "營業人名稱（蓋公司印鑑章）" in stripped or "立承諾書營業人" in stripped or "立切結書營業人" in stripped):
            if "此致" in stripped:
                p_to = doc.add_paragraph()
                p_to.paragraph_format.space_before = Pt(14)
                p_to.paragraph_format.space_after = Pt(2)
                run_to = p_to.add_run("此致")
                set_run_font(run_to, size_pt=11, bold=True)
                i += 1
                if i < len(lines):
                    next_l = lines[i].strip()
                    if next_l.startswith("**") and next_l.endswith("**"):
                        p_agency = doc.add_paragraph()
                        p_agency.paragraph_format.space_before = Pt(2)
                        p_agency.paragraph_format.space_after = Pt(8)
                        run_ag = p_agency.add_run(next_l.strip('*'))
                        set_run_font(run_ag, size_pt=12, bold=True, color=COLOR_NAVY)
                        i += 1
            add_seal_box(doc, has_rep_id=has_rep_id)
            seal_box_added = True
            # Skip any leftover signature text lines in markdown
            while i < len(lines):
                i += 1
            continue
            
        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        parse_formatted_text(p, stripped, base_size=10.5)
        i += 1
        
    # Final check if statutory form hasn't had seal box added yet
    if is_statutory_form and not seal_box_added:
        add_seal_box(doc, has_rep_id=has_rep_id)
        seal_box_added = True
        
    doc.save(docx_path)
    print(f"Generated: {docx_path}")

def main():
    md_files = [f for f in os.listdir(DOCS_DIR) if f.endswith(".md")]
    md_files.sort()
    
    print(f"Found {len(md_files)} markdown files in {DOCS_DIR}")
    for md_file in md_files:
        md_path = os.path.join(DOCS_DIR, md_file)
        docx_file = os.path.splitext(md_file)[0] + ".docx"
        docx_path = os.path.join(DOCS_DIR, docx_file)
        convert_md_file(md_path, docx_path)

if __name__ == "__main__":
    main()
